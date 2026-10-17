
from io import BytesIO
from pathlib import Path

from docx import Document
from pypdf import PdfReader


MAX_FILE_SIZE = 10 * 1024 * 1024  # 10 MB
SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt"}


def extract_text(filename: str, file_bytes: bytes) -> str:
    """Extract readable text from a PDF, DOCX, or TXT file."""

    if not isinstance(filename, str) or not filename.strip():
        raise ValueError("A valid filename is required.")

    if not isinstance(file_bytes, bytes):
        raise ValueError("File content must be provided as bytes.")

    extension = Path(filename).suffix.lower()

    if extension not in SUPPORTED_EXTENSIONS:
        raise ValueError(
            "Unsupported file type. Please upload a PDF, DOCX, or TXT file."
        )

    if not file_bytes:
        raise ValueError("The uploaded file is empty.")

    if len(file_bytes) > MAX_FILE_SIZE:
        raise ValueError("The file is too large. Maximum size is 10 MB.")

    try:
        if extension == ".txt":
            text = file_bytes.decode("utf-8-sig")

        elif extension == ".pdf":
            reader = PdfReader(BytesIO(file_bytes))

            if reader.is_encrypted:
                raise ValueError(
                    "This PDF is password-protected. Upload an unlocked PDF."
                )

            text = "\n".join(
                page.extract_text() or "" for page in reader.pages
            )

        elif extension == ".docx":
            document = Document(BytesIO(file_bytes))
            paragraphs = [p.text for p in document.paragraphs]
            table_text = [
                cell.text
                for table in document.tables
                for row in table.rows
                for cell in row.cells
            ]
            text = "\n".join(paragraphs + table_text)

    except ValueError:
        raise
    except Exception as exc:
        raise ValueError(
            f"Could not read this {extension[1:].upper()} file. "
            "Please check that it is a valid, uncorrupted document."
        ) from exc

    cleaned_text = "\n".join(
        line.strip() for line in text.splitlines() if line.strip()
    )

    if not cleaned_text:
        raise ValueError(
            "No readable text was found. Scanned PDFs may require OCR, "
            "which is not supported by this utility."
        )

    return cleaned_text
